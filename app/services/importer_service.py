import os
import re
import csv
from typing import List, Dict, Any, Optional
import openpyxl
from sqlalchemy.orm import Session
from app.db.models import Pelanggan

def clean_string(val: Any) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None

def clean_id_pelanggan(val: Any, fallback_idx: int) -> str:
    if val is None or str(val).strip() == "":
        return f"CUST-{fallback_idx:04d}"
    s = str(val).strip()
    # Jika notasi eksponensial dalam float seperti 2.61E+11
    if "E+" in s or "e+" in s:
        try:
            return str(int(float(s)))
        except ValueError:
            pass
    if s.endswith(".0"):
        s = s[:-2]
    return s

def clean_float(val: Any) -> Optional[float]:
    if val is None:
        return None
    s = str(val).replace("dBm", "").replace("dbm", "").replace(",", ".").strip()
    m = re.search(r"[-+]?\d*\.?\d+", s)
    if m:
        try:
            return float(m.group(0))
        except ValueError:
            return None
    return None

class ImporterService:

    @staticmethod
    def import_from_csv(file_path: str, db: Session) -> int:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File {file_path} tidak ditemukan")

        imported_count = 0
        seen_ids = set()

        # Ambil id yang sudah ada di database
        for (existing_id,) in db.query(Pelanggan.id_pelanggan).all():
            seen_ids.add(str(existing_id))

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            for idx, row in enumerate(reader, start=1):
                nama = clean_string(row.get("Nama"))
                ip = clean_string(row.get("IP Router"))
                if not nama or not ip:
                    continue

                raw_id = clean_id_pelanggan(row.get("ID Pelanggan"), idx)
                id_pel = raw_id
                if id_pel in seen_ids:
                    id_pel = f"{raw_id}_{idx}"
                seen_ids.add(id_pel)

                modem = clean_string(row.get("Jenis Modem")) or "GM220-S"
                pop = clean_string(row.get("POP")) or "Server Cabang"
                redaman = clean_float(row.get("Redaman"))

                new_pelanggan = Pelanggan(
                    id_pelanggan=id_pel,
                    nama=nama,
                    alamat=clean_string(row.get("Alamat")),
                    no_hp=clean_string(row.get("No HP")),
                    pop=pop,
                    ip_router=ip,
                    paket=clean_string(row.get("Paket")),
                    jenis_modem=modem,
                    mac_address=clean_string(row.get("MAC Address")),
                    redaman_baseline=redaman,
                    nama_wifi=clean_string(row.get("Nama Wifi")),
                    password_wifi=clean_string(row.get("Password wifi")),
                    user_admin=clean_string(row.get("USER ADMIN")),
                    pass_admin=clean_string(row.get("PASS ADMIN")),
                    snmp_community="public",
                    is_active=True
                )
                db.add(new_pelanggan)
                imported_count += 1

        db.commit()
        return imported_count

    @staticmethod
    def import_from_excel(file_path: str, db: Session, sheet_names: Optional[List[str]] = None) -> int:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File {file_path} tidak ditemukan")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        target_sheets = sheet_names or ["PELANGGAN CABANG"]
        imported_count = 0
        global_idx = 1

        for sname in target_sheets:
            if sname not in wb.sheetnames:
                continue
            ws = wb[sname]
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                continue

            headers = [str(c).strip() if c is not None else "" for c in rows[0]]
            col_map = {h.lower(): i for i, h in enumerate(headers) if h}

            for row in rows[1:]:
                def get_val(key_fragment: str):
                    for h, i in col_map.items():
                        if key_fragment in h:
                            return row[i] if i < len(row) else None
                    return None

                nama = clean_string(get_val("nama"))
                ip = clean_string(get_val("ip router") or get_val("ip remote") or get_val("ip"))
                if not nama or not ip:
                    continue

                id_pel = clean_id_pelanggan(get_val("id pelanggan"), global_idx)
                global_idx += 1
                modem = clean_string(get_val("jenis modem") or get_val("modem")) or "GM220-S"
                pop = clean_string(get_val("pop")) or sname.replace("PELANGGAN ", "Server ")
                redaman = clean_float(get_val("redaman"))

                existing = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == id_pel).first()
                if existing:
                    existing.nama = nama
                    existing.alamat = clean_string(get_val("alamat"))
                    existing.no_hp = clean_string(get_val("no hp") or get_val("nomor wa"))
                    existing.pop = pop
                    existing.ip_router = ip
                    existing.paket = clean_string(get_val("paket"))
                    existing.jenis_modem = modem
                    existing.mac_address = clean_string(get_val("mac"))
                    existing.redaman_baseline = redaman
                    existing.nama_wifi = clean_string(get_val("nama wifi") or get_val("wifi"))
                    existing.password_wifi = clean_string(get_val("password wifi") or get_val("pswd"))
                    existing.user_admin = clean_string(get_val("user admin"))
                    existing.pass_admin = clean_string(get_val("pass admin") or get_val("passwrod admin"))
                else:
                    new_pelanggan = Pelanggan(
                        id_pelanggan=id_pel,
                        nama=nama,
                        alamat=clean_string(get_val("alamat")),
                        no_hp=clean_string(get_val("no hp") or get_val("nomor wa")),
                        pop=pop,
                        ip_router=ip,
                        paket=clean_string(get_val("paket")),
                        jenis_modem=modem,
                        mac_address=clean_string(get_val("mac")),
                        redaman_baseline=redaman,
                        nama_wifi=clean_string(get_val("nama wifi") or get_val("wifi")),
                        password_wifi=clean_string(get_val("password wifi") or get_val("pswd")),
                        user_admin=clean_string(get_val("user admin")),
                        pass_admin=clean_string(get_val("pass admin") or get_val("passwrod admin")),
                        snmp_community="public",
                        is_active=True
                    )
                    db.add(new_pelanggan)
                imported_count += 1

        db.commit()
        return imported_count
