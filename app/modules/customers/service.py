import uuid
import csv
import io
import re
from datetime import datetime, timedelta
from typing import Optional, List, Tuple, Dict, Any
from sqlalchemy.orm import Session
from sqlalchemy import or_, func

import openpyxl
import pandas as pd
import json

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.db.models import Pelanggan, LogPerformaONT
from app.modules.customers.schemas import CustomerCreate, CustomerUpdate

class CustomerService:

    @staticmethod
    def read_csv_safely(file_content: bytes) -> pd.DataFrame:
        """
        Membaca file CSV dengan multi-encoding (utf-8-sig, utf-8, cp1252, latin-1)
        dan auto-detect delimiter (; , \\t |), serta membersihkan BOM header.
        """
        encodings_to_try = ['utf-8-sig', 'utf-8', 'cp1252', 'latin-1', 'iso-8859-1']
        decoded_text = None
        for enc in encodings_to_try:
            try:
                decoded_text = file_content.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        if decoded_text is None:
            decoded_text = file_content.decode('utf-8', errors='ignore')

        # Deteksi delimiter terbaik berdasarkan baris pertama data
        sample = decoded_text[:8192]
        lines = [l for l in sample.splitlines() if l.strip()]
        first_line = lines[0] if lines else ""

        candidates = [';', ',', '\t', '|']
        counts = {c: first_line.count(c) for c in candidates}
        best_sep = max(counts, key=counts.get)
        if counts[best_sep] == 0:
            try:
                dialect = csv.Sniffer().sniff(sample)
                best_sep = dialect.delimiter
            except Exception:
                best_sep = ','

        df = pd.read_csv(io.StringIO(decoded_text), sep=best_sep, dtype=str)
        # Bersihkan BOM, spasi, dan quotes tersembunyi pada header
        df.columns = [str(c).replace('\ufeff', '').strip().strip('"').strip("'") for c in df.columns]
        return df

    @staticmethod
    def generate_customer_id_from_ip(ip: str, sequence: int = 1) -> str:
        """
        Menghasilkan ID pelanggan dengan format P(idrouter)00001
        Contoh: IP 10.10.7.62 -> P101076200001
        Jika IP kosong -> PLG00001
        """
        clean_ip = re.sub(r'[^0-9]', '', str(ip or ''))
        seq_str = f"{sequence:05d}"
        if clean_ip:
            return f"P{clean_ip}{seq_str}"
        return f"PLG{seq_str}"

    @staticmethod
    def analyze_import_file(file_content: bytes, filename: str) -> dict:
        filename_lower = filename.lower()
        if filename_lower.endswith('.csv'):
            try:
                df = CustomerService.read_csv_safely(file_content)
                cols = [str(c).strip() for c in df.columns if str(c).strip() and not str(c).startswith('Unnamed:')]
                
                return {
                    "is_excel": False,
                    "sheets": [{
                        "name": "CSV Data",
                        "columns": cols,
                        "header_row": 0,
                        "row_count": len(df)
                    }]
                }
            except Exception as e:
                raise ValueError(f"Gagal membaca file CSV: {str(e)}")
        elif filename_lower.endswith(('.xls', '.xlsx')):
            try:
                excel = pd.ExcelFile(io.BytesIO(file_content))
                sheets_data = []
                for sheet in excel.sheet_names:
                    # Deteksi baris header secara cerdas (periksa 10 baris pertama)
                    df_raw = pd.read_excel(excel, sheet_name=sheet, header=None, nrows=10)
                    if df_raw.empty:
                        sheets_data.append({
                            "name": sheet,
                            "columns": [],
                            "header_row": 0,
                            "row_count": 0
                        })
                        continue

                    best_row = 0
                    max_non_null = 0
                    for r in range(min(len(df_raw), 8)):
                        row_vals = [
                            str(x).strip() for x in df_raw.iloc[r] 
                            if pd.notna(x) and str(x).strip() != '' and not str(x).startswith('Unnamed:')
                        ]
                        if len(row_vals) > max_non_null:
                            max_non_null = len(row_vals)
                            best_row = r

                    cols = [
                        str(x).strip() for x in df_raw.iloc[best_row] 
                        if pd.notna(x) and str(x).strip() != '' and not str(x).startswith('Unnamed:')
                    ]

                    # Hitung estimasi baris data
                    df_full = pd.read_excel(excel, sheet_name=sheet, header=best_row)
                    row_count = len(df_full)

                    sheets_data.append({
                        "name": sheet,
                        "columns": cols,
                        "header_row": best_row,
                        "row_count": row_count
                    })

                return {
                    "is_excel": True,
                    "sheets": sheets_data
                }
            except Exception as e:
                raise ValueError(f"Gagal membaca file Excel: {str(e)}")
        else:
            raise ValueError("Format file tidak didukung. Harap gunakan berkas .xlsx, .xls, atau .csv")

    @staticmethod
    def preview_import_file(file_content: bytes, filename: str, sheet_name: str, mapping: dict, db: Session) -> dict:
        import math
        from app.db.models import Pelanggan

        filename_lower = filename.lower()
        try:
            if filename_lower.endswith('.csv'):
                df = CustomerService.read_csv_safely(file_content)
            else:
                excel = pd.ExcelFile(io.BytesIO(file_content))
                target_sheet = sheet_name if sheet_name in excel.sheet_names else excel.sheet_names[0]
                df_raw = pd.read_excel(excel, sheet_name=target_sheet, header=None, nrows=10)
                best_row = 0
                max_non_null = 0
                for r in range(min(len(df_raw), 8)):
                    row_vals = [
                        str(x).strip() for x in df_raw.iloc[r] 
                        if pd.notna(x) and str(x).strip() != '' and not str(x).startswith('Unnamed:')
                    ]
                    if len(row_vals) > max_non_null:
                        max_non_null = len(row_vals)
                        best_row = r

                df = pd.read_excel(excel, sheet_name=target_sheet, header=best_row)

            # Normalisasi kolom
            df.columns = [str(c).strip() for c in df.columns]
            df = df.fillna("")

            # Query data existing di database untuk deteksi ID dan IP unik
            existing_customers = db.query(Pelanggan.id_pelanggan, Pelanggan.ip_router, Pelanggan.nama).filter(Pelanggan.is_active == True).all()
            existing_ids = {row[0] for row in existing_customers if row[0]}
            existing_ips = {row[1]: (row[0], row[2]) for row in existing_customers if row[1]}

            seen_ids_in_batch = {} # id_pelanggan -> original_row
            seen_ips_in_batch = {}  # ip_router -> original_row
            ip_sequence_counter = {} # clean_ip -> sequence

            mapped_data = []

            for index, row in df.iterrows():
                row_dict = row.to_dict()
                mapped_row = {}
                has_any_data = False

                for db_field, excel_col in mapping.items():
                    if excel_col and excel_col in row_dict:
                        val = row_dict[excel_col]
                        if isinstance(val, float) and math.isnan(val):
                            val = ""
                        val = str(val).strip()
                        if val.endswith('.0') and db_field in ['id_pelanggan', 'no_hp', 'ip_router']:
                            val = val[:-2]
                        mapped_row[db_field] = val
                        if val:
                            has_any_data = True
                    else:
                        mapped_row[db_field] = ""

                # Lewati baris yang benar-benar kosong
                if not has_any_data:
                    continue

                nama = mapped_row.get('nama', '').strip()
                ip_router = mapped_row.get('ip_router', '').strip()

                # Cek ID pelanggan: Jika kosong, buat format P(idrouter)00001
                id_pel = mapped_row.get('id_pelanggan', '').strip()
                if not id_pel:
                    clean_ip = re.sub(r'[^0-9]', '', ip_router)
                    ip_key = clean_ip if clean_ip else "PLG"
                    ip_sequence_counter[ip_key] = ip_sequence_counter.get(ip_key, 0) + 1
                    id_pel = CustomerService.generate_customer_id_from_ip(ip_router, ip_sequence_counter[ip_key])
                    mapped_row['id_pelanggan'] = id_pel

                _status = "valid"
                _message = "Data baru siap di-import"

                if not nama or not ip_router:
                    _status = "error"
                    _message = "Nama Pelanggan dan IP Router wajib diisi"
                # 1. Pengecekan Duplikat IP Router (Prioritas Utama)
                elif ip_router in existing_ips:
                    _status = "duplicate"
                    ex_id, ex_name = existing_ips[ip_router]
                    _message = f"IP Router '{ip_router}' sudah terdaftar di database (Pelanggan: {ex_name} [{ex_id}])"
                elif ip_router in seen_ips_in_batch:
                    _status = "duplicate"
                    _message = f"IP Router '{ip_router}' duplikat di berkas ini (sama dengan baris {seen_ips_in_batch[ip_router]})"
                # 2. Pengecekan Duplikat ID Pelanggan
                elif id_pel in existing_ids:
                    _status = "duplicate"
                    _message = f"ID Pelanggan '{id_pel}' sudah terdaftar di database"
                elif id_pel in seen_ids_in_batch:
                    _status = "duplicate"
                    _message = f"ID Pelanggan '{id_pel}' duplikat di berkas ini (sama dengan baris {seen_ids_in_batch[id_pel]})"
                else:
                    seen_ids_in_batch[id_pel] = index + 1
                    seen_ips_in_batch[ip_router] = index + 1

                # Catat ke seen_batch jika duplikat agar baris berikutnya juga terdeteksi
                if ip_router and ip_router not in seen_ips_in_batch:
                    seen_ips_in_batch[ip_router] = index + 1
                if id_pel and id_pel not in seen_ids_in_batch:
                    seen_ids_in_batch[id_pel] = index + 1

                mapped_data.append({
                    "original_row": index + 1,
                    "data": mapped_row,
                    "_status": _status,
                    "_action": "update" if _status == "duplicate" else "insert",
                    "_message": _message
                })

            valid_count = sum(1 for r in mapped_data if r["_status"] == "valid")
            dup_count = sum(1 for r in mapped_data if r["_status"] == "duplicate")
            err_count = sum(1 for r in mapped_data if r["_status"] == "error")

            return {
                "total_rows": len(mapped_data),
                "valid_count": valid_count,
                "duplicates_count": dup_count,
                "errors_count": err_count,
                "preview_data": mapped_data,
                "existing_ids": list(existing_ids),
                "existing_ips": list(existing_ips.keys())
            }
        except Exception as e:
            raise ValueError(f"Gagal memproses file untuk preview: {str(e)}")

    @staticmethod
    def execute_json_import(db: Session, data_list: list) -> dict:
        from app.db.models import Pelanggan

        # 1. Validasi Ketat: Tolak eksekusi jika terdapat duplikasi IP Router atau ID Pelanggan
        active_items = [item for item in data_list if item.get("_action") != "skip" and item.get("_status") != "error"]
        
        batch_ids = []
        batch_ips = []
        duplicate_reasons = []

        # Ambil IP yang sudah aktif di DB untuk verifikasi tabrakan IP
        db_ips = {row[0]: row[1] for row in db.query(Pelanggan.ip_router, Pelanggan.id_pelanggan).filter(Pelanggan.is_active == True).all() if row[0]}

        for item in active_items:
            row_data = item.get("data", item)
            id_pel = str(row_data.get("id_pelanggan", "")).strip()
            ip_r = str(row_data.get("ip_router", "")).strip()

            if id_pel:
                if id_pel in batch_ids:
                    duplicate_reasons.append(f"ID Pelanggan ganda di dalam data: '{id_pel}'")
                batch_ids.append(id_pel)

            if ip_r:
                if ip_r in batch_ips:
                    duplicate_reasons.append(f"IP Router ganda di dalam data: '{ip_r}'")
                batch_ips.append(ip_r)
                # Jika baris ini baru (insert) tapi IP-nya sudah dipakai ID lain di DB
                if ip_r in db_ips and db_ips[ip_r] != id_pel:
                    duplicate_reasons.append(f"IP Router '{ip_r}' sudah digunakan oleh pelanggan lain di database (ID: {db_ips[ip_r]})")

        if duplicate_reasons:
            reasons_summary = "; ".join(duplicate_reasons[:3])
            if len(duplicate_reasons) > 3:
                reasons_summary += f" ... dan {len(duplicate_reasons) - 3} lainnya"
            raise ValueError(f"Eksekusi Ditolak: Ditemukan duplikasi ({reasons_summary}). Seluruh ID Pelanggan dan IP Router yang bertabrakan WAJIB dibenarkan terlebih dahulu sebelum import dapat dieksekusi!")

        imported_count = 0
        updated_count = 0
        skipped_count = 0
        failed_count = 0
        errors = []

        for item in data_list:
            try:
                # Dukung format payload langsung maupun terbungkus {data: ...}
                row_data = item.get("data", item)
                action = item.get("_action", "insert")
                status = item.get("_status", "valid")

                if action == "skip" or status == "error":
                    skipped_count += 1
                    continue

                id_pel = str(row_data.get("id_pelanggan", "")).strip()
                nama = str(row_data.get("nama", "")).strip()
                ip_router = str(row_data.get("ip_router", "")).strip()

                if not id_pel or not nama or not ip_router:
                    failed_count += 1
                    errors.append(f"Baris tidak lengkap: ID='{id_pel}', Nama='{nama}', IP='{ip_router}'")
                    continue

                existing = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == id_pel).first()
                if existing:
                    existing.nama = nama
                    if row_data.get("alamat"): existing.alamat = str(row_data["alamat"]).strip()
                    if row_data.get("no_hp"): existing.no_hp = str(row_data["no_hp"]).strip()
                    if row_data.get("pop"): existing.pop = str(row_data["pop"]).strip()
                    existing.ip_router = ip_router
                    if row_data.get("paket"): existing.paket = str(row_data["paket"]).strip()
                    if row_data.get("jenis_modem"): existing.jenis_modem = str(row_data["jenis_modem"]).strip()
                    if row_data.get("mac_address"): existing.mac_address = str(row_data["mac_address"]).strip()
                    
                    if row_data.get("redaman_baseline"):
                        try:
                            existing.redaman_baseline = float(row_data["redaman_baseline"])
                        except (ValueError, TypeError):
                            pass
                            
                    if row_data.get("nama_wifi"): existing.nama_wifi = str(row_data["nama_wifi"]).strip()
                    if row_data.get("password_wifi"): existing.password_wifi = str(row_data["password_wifi"]).strip()
                    if row_data.get("user_admin"): existing.user_admin = str(row_data["user_admin"]).strip()
                    if row_data.get("pass_admin"): existing.pass_admin = str(row_data["pass_admin"]).strip()
                    existing.status_kredensial = "UNTESTED"
                    existing.is_active = True
                    
                    updated_count += 1
                else:
                    redaman_baseline = None
                    if row_data.get("redaman_baseline"):
                        try:
                            redaman_baseline = float(row_data["redaman_baseline"])
                        except (ValueError, TypeError):
                            pass
                            
                    new_pelanggan = Pelanggan(
                        id_pelanggan=id_pel,
                        nama=nama,
                        alamat=str(row_data.get("alamat", "")).strip() or None,
                        no_hp=str(row_data.get("no_hp", "")).strip() or None,
                        pop=str(row_data.get("pop", "Server Cabang")).strip() or "Server Cabang",
                        ip_router=ip_router,
                        paket=str(row_data.get("paket", "")).strip() or None,
                        jenis_modem=str(row_data.get("jenis_modem", "GM220-S")).strip() or "GM220-S",
                        mac_address=str(row_data.get("mac_address", "")).strip() or None,
                        redaman_baseline=redaman_baseline,
                        nama_wifi=str(row_data.get("nama_wifi", "")).strip() or None,
                        password_wifi=str(row_data.get("password_wifi", "")).strip() or None,
                        user_admin=str(row_data.get("user_admin", "admin")).strip() or "admin",
                        pass_admin=str(row_data.get("pass_admin", "")).strip() or None,
                        status_kredensial="UNTESTED",
                        is_active=True
                    )
                    db.add(new_pelanggan)
                    imported_count += 1
            except Exception as e:
                failed_count += 1
                errors.append(f"Gagal memproses ID {row_data.get('id_pelanggan', 'Unknown')}: {str(e)}")

        db.commit()

        return {
            "total_processed": imported_count + updated_count + skipped_count + failed_count,
            "imported": imported_count,
            "updated": updated_count,
            "skipped": skipped_count,
            "failed": failed_count,
            "errors": errors[:5]
        }

    @staticmethod
    def get_customers(
        db: Session,
        q: Optional[str] = None,
        pop: Optional[str] = None,
        status: Optional[str] = None,
        monitoring: Optional[str] = None,
        range_type: Optional[str] = "all",
        start_date: Optional[str] = None,
        end_date: Optional[str] = None,
        sort_by: Optional[str] = "id",
        sort_dir: str = "asc",
        page: int = 1,
        limit: int = 25
    ) -> Tuple[int, List[Dict[str, Any]]]:
        # Subquery ID log terbaru per id_pelanggan
        subq = (
            db.query(
                LogPerformaONT.id_pelanggan,
                func.max(LogPerformaONT.id).label('max_log_id')
            )
            .group_by(LogPerformaONT.id_pelanggan)
            .subquery()
        )

        query = (
            db.query(Pelanggan, LogPerformaONT)
            .outerjoin(subq, subq.c.id_pelanggan == Pelanggan.id_pelanggan)
            .outerjoin(LogPerformaONT, LogPerformaONT.id == subq.c.max_log_id)
            .filter(Pelanggan.is_active == True)
        )

        if q:
            search_pattern = f"%{q.strip()}%"
            query = query.filter(
                or_(
                    Pelanggan.nama.ilike(search_pattern),
                    Pelanggan.ip_router.ilike(search_pattern),
                    Pelanggan.mac_address.ilike(search_pattern),
                    Pelanggan.id_pelanggan.ilike(search_pattern)
                )
            )

        if pop and pop != "Semua POP":
            query = query.filter(Pelanggan.pop == pop)

        if status and status != "Semua Status":
            if status == "NORMAL":
                query = query.filter(or_(LogPerformaONT.status_koneksi == "NORMAL", LogPerformaONT.status_koneksi == None))
            elif status == "WARNING":
                query = query.filter(LogPerformaONT.status_koneksi == "WARNING")
            elif status in ("CRITICAL", "LOS"):
                query = query.filter(or_(LogPerformaONT.status_koneksi == "CRITICAL", LogPerformaONT.status_koneksi == "LOS"))
            elif status in ("NONAKTIF", "OFF"):
                query = query.filter(Pelanggan.is_monitored == False)
            else:
                query = query.filter(LogPerformaONT.status_koneksi == status)

        if monitoring and monitoring != "all":
            if monitoring in ("active", "true", "1"):
                query = query.filter(Pelanggan.is_monitored == True)
            elif monitoring in ("inactive", "false", "0"):
                query = query.filter(Pelanggan.is_monitored == False)

        # Filter Rentang Waktu Tanggal (Berdasarkan waktu cek log terakhir atau tanggal terdaftar)
        now = datetime.utcnow()
        date_col = func.coalesce(LogPerformaONT.waktu_cek, Pelanggan.created_at)
        if range_type == "today":
            cutoff = datetime(now.year, now.month, now.day)
            query = query.filter(date_col >= cutoff)
        elif range_type == "7d":
            cutoff = now - timedelta(days=7)
            query = query.filter(date_col >= cutoff)
        elif range_type == "30d":
            cutoff = now - timedelta(days=30)
            query = query.filter(date_col >= cutoff)
        elif range_type == "custom" and start_date and end_date:
            try:
                s_dt = datetime.strptime(start_date.strip(), "%Y-%m-%d")
                e_dt = datetime.strptime(end_date.strip(), "%Y-%m-%d") + timedelta(days=1)
                query = query.filter(date_col >= s_dt, date_col < e_dt)
            except ValueError:
                pass

        total_count = query.count()

        # Dynamic sorting
        col_map = {
            "id_pelanggan": Pelanggan.id_pelanggan,
            "nama": Pelanggan.nama,
            "pop": Pelanggan.pop,
            "ip_router": Pelanggan.ip_router,
            "paket": Pelanggan.paket,
            "jenis_modem": Pelanggan.jenis_modem,
            "redaman_baseline": Pelanggan.redaman_baseline,
            "status_kredensial": Pelanggan.status_kredensial,
            "redaman_current": LogPerformaONT.rx_power,
            "status": LogPerformaONT.status_koneksi,
            "is_monitored": Pelanggan.is_monitored,
            "waktu_cek": LogPerformaONT.waktu_cek,
            "created_at": Pelanggan.created_at,
            "id": Pelanggan.id
        }
        target_col = col_map.get(sort_by, Pelanggan.id)
        if sort_dir.lower() == "desc":
            query = query.order_by(target_col.desc())
        else:
            query = query.order_by(target_col.asc())

        offset = (page - 1) * limit
        rows = query.offset(offset).limit(limit).all()

        result = []
        for c, latest_log in rows:
            current_status = latest_log.status_koneksi if latest_log else "NORMAL"
            current_rx = float(latest_log.rx_power) if (latest_log and latest_log.rx_power is not None) else (float(c.redaman_baseline) if c.redaman_baseline else None)
            last_check = latest_log.waktu_cek.strftime("%d/%m %H:%M") if latest_log else "-"

            result.append({
                "id_pelanggan": c.id_pelanggan,
                "nama": c.nama,
                "alamat": c.alamat or "-",
                "no_hp": c.no_hp or "-",
                "pop": c.pop,
                "ip_router": c.ip_router,
                "paket": c.paket or "-",
                "jenis_modem": c.jenis_modem,
                "mac_address": c.mac_address or "-",
                "redaman_baseline": float(c.redaman_baseline) if c.redaman_baseline else None,
                "redaman_current": current_rx,
                "status": current_status,
                "last_check": last_check,
                "nama_wifi": c.nama_wifi or "-",
                "password_wifi": c.password_wifi or "-",
                "user_admin": c.user_admin or "-",
                "pass_admin": c.pass_admin or "-",
                "status_kredensial": getattr(c, "status_kredensial", "UNTESTED") or "UNTESTED",
                "is_monitored": bool(getattr(c, "is_monitored", True))
            })

        return total_count, result

    @staticmethod
    def toggle_monitoring(db: Session, id_pelanggan: str, is_monitored: Optional[bool] = None) -> Optional[Pelanggan]:
        cust = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == id_pelanggan).first()
        if not cust:
            return None
        if is_monitored is not None:
            cust.is_monitored = bool(is_monitored)
        else:
            cust.is_monitored = not bool(getattr(cust, "is_monitored", True))
        db.commit()
        db.refresh(cust)
        return cust

    @staticmethod
    def get_customer_detail(db: Session, id_pelanggan: str) -> Optional[Tuple[Pelanggan, List[LogPerformaONT]]]:
        cust = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == id_pelanggan).first()
        if not cust:
            return None

        logs = (
            db.query(LogPerformaONT)
            .filter(LogPerformaONT.id_pelanggan == id_pelanggan)
            .order_by(LogPerformaONT.waktu_cek.desc())
            .limit(50)
            .all()
        )
        return cust, logs

    @staticmethod
    def get_min_date(db: Session) -> str:
        earliest_log = db.query(func.min(LogPerformaONT.waktu_cek)).scalar()
        earliest_cust = db.query(func.min(Pelanggan.created_at)).scalar()
        dates = [d for d in [earliest_log, earliest_cust] if d is not None]
        if dates:
            return min(dates).strftime("%Y-%m-%d")
        return "2026-09-01"

    @staticmethod
    def get_pop_list(db: Session) -> List[str]:
        rows = db.query(Pelanggan.pop).filter(Pelanggan.is_active == True).distinct().all()
        return [r[0] for r in rows if r[0]]

    @staticmethod
    def get_customer_stats(db: Session) -> Dict[str, Any]:
        # Hanya hitung pelanggan yang aktif dan dipantau (is_monitored == True) di card
        monitored_active = db.query(Pelanggan).filter(
            Pelanggan.is_active == True,
            Pelanggan.is_monitored == True
        ).count()
        monitored_inactive = db.query(Pelanggan).filter(
            Pelanggan.is_active == True,
            Pelanggan.is_monitored == False
        ).count()
        total_all = db.query(Pelanggan).filter(Pelanggan.is_active == True).count()

        total_pelanggan = monitored_active  # Nilai card dikurangi pelanggan yang OFF

        pop_counts = (
            db.query(Pelanggan.pop, func.count(Pelanggan.id))
            .filter(
                Pelanggan.is_active == True,
                Pelanggan.is_monitored == True
            )
            .group_by(Pelanggan.pop)
            .all()
        )

        subq = (
            db.query(
                LogPerformaONT.id_pelanggan,
                func.max(LogPerformaONT.id).label('max_log_id')
            )
            .join(Pelanggan, Pelanggan.id_pelanggan == LogPerformaONT.id_pelanggan)
            .filter(
                Pelanggan.is_active == True,
                Pelanggan.is_monitored == True
            )
            .group_by(LogPerformaONT.id_pelanggan)
            .subquery()
        )

        status_counts = (
            db.query(
                LogPerformaONT.status_koneksi,
                func.count(Pelanggan.id)
            )
            .join(subq, subq.c.id_pelanggan == Pelanggan.id_pelanggan)
            .join(LogPerformaONT, LogPerformaONT.id == subq.c.max_log_id)
            .filter(
                Pelanggan.is_active == True,
                Pelanggan.is_monitored == True
            )
            .group_by(LogPerformaONT.status_koneksi)
            .all()
        )
        status_dict = {s[0]: s[1] for s in status_counts}
        normal_count = status_dict.get("NORMAL", 0)
        warning_count = status_dict.get("WARNING", 0)
        critical_count = status_dict.get("CRITICAL", 0)
        los_count = status_dict.get("LOS", 0)

        total_with_logs = sum(status_dict.values())
        if total_pelanggan > total_with_logs:
            normal_count += (total_pelanggan - total_with_logs)

        return {
            "total": total_pelanggan,
            "total_all": total_all,
            "normal": normal_count,
            "warning": warning_count,
            "critical": critical_count,
            "los": los_count,
            "critical_los": critical_count + los_count,
            "monitored_active": monitored_active,
            "monitored_inactive": monitored_inactive,
            "pop_distribution": [{"pop": p[0], "count": p[1]} for p in pop_counts]
        }

    @staticmethod
    def create_customer(db: Session, data: CustomerCreate) -> Pelanggan:
        if not data.id_pelanggan or not str(data.id_pelanggan).strip():
            # Cari urutan sequence untuk IP router terkait jika sudah ada di DB
            clean_ip = "".join(filter(str.isdigit, str(data.ip_router or "0")))
            existing_count = db.query(Pelanggan).filter(Pelanggan.id_pelanggan.like(f"P{clean_ip}%")).count()
            data.id_pelanggan = CustomerService.generate_customer_id_from_ip(data.ip_router or "10.10.0.1", existing_count + 1)

        # Validasi duplikasi ID Pelanggan
        existing_id = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == data.id_pelanggan).first()
        if existing_id:
            raise ValueError(f"ID Pelanggan '{data.id_pelanggan}' sudah terdaftar")

        # Validasi duplikasi IP Router aktif
        existing_ip = db.query(Pelanggan).filter(
            Pelanggan.ip_router == data.ip_router,
            Pelanggan.is_active == True
        ).first()
        if existing_ip:
            raise ValueError(f"IP Router '{data.ip_router}' sudah digunakan oleh pelanggan '{existing_ip.nama}'")

        new_cust = Pelanggan(
            id_pelanggan=data.id_pelanggan,
            nama=data.nama,
            alamat=data.alamat,
            no_hp=data.no_hp,
            pop=data.pop,
            ip_router=data.ip_router,
            paket=data.paket,
            jenis_modem=data.jenis_modem,
            mac_address=data.mac_address,
            redaman_baseline=data.redaman_baseline,
            nama_wifi=data.nama_wifi,
            password_wifi=data.password_wifi,
            user_admin=data.user_admin,
            pass_admin=data.pass_admin,
            snmp_community=data.snmp_community,
            is_monitored=getattr(data, "is_monitored", True) if getattr(data, "is_monitored", None) is not None else True,
            is_active=True
        )
        db.add(new_cust)
        db.commit()
        db.refresh(new_cust)
        return new_cust

    @staticmethod
    def update_customer(db: Session, id_pelanggan: str, data: CustomerUpdate) -> Optional[Pelanggan]:
        cust = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == id_pelanggan).first()
        if not cust:
            return None

        update_dict = data.dict(exclude_unset=True)

        if "ip_router" in update_dict and update_dict["ip_router"]:
            new_ip = str(update_dict["ip_router"]).strip()
            # Periksa apakah IP sudah digunakan oleh pelanggan aktif lain
            existing_ip = db.query(Pelanggan).filter(
                Pelanggan.ip_router == new_ip,
                Pelanggan.is_active == True,
                Pelanggan.id_pelanggan != id_pelanggan
            ).first()
            if existing_ip:
                raise ValueError(f"IP Router '{new_ip}' sudah digunakan oleh pelanggan '{existing_ip.nama}' (ID: {existing_ip.id_pelanggan})")

        for field, val in update_dict.items():
            if hasattr(cust, field) and val is not None:
                setattr(cust, field, val)

        db.commit()
        db.refresh(cust)
        return cust

    @staticmethod
    def delete_customer(db: Session, id_pelanggan: str) -> bool:
        cust = db.query(Pelanggan).filter(Pelanggan.id_pelanggan == id_pelanggan).first()
        if not cust:
            return False

        # Soft delete
        cust.is_active = False
        db.commit()
        return True

    @staticmethod
    def bulk_delete_customers(db: Session, id_list: List[str]) -> int:
        """Soft delete banyak pelanggan sekaligus berdasarkan list id_pelanggan"""
        if not id_list:
            return 0

        affected = (
            db.query(Pelanggan)
            .filter(Pelanggan.id_pelanggan.in_(id_list))
            .update({Pelanggan.is_active: False}, synchronize_session=False)
        )
        db.commit()
        return affected

    @staticmethod
    def import_customers_from_csv(db: Session, file_content: bytes) -> Dict[str, Any]:
        """
        Membaca dan memproses isi file CSV pelanggan dengan Smart Header Detection.
        Mendukung berbagai nama header sinonim tanpa mewajibkan template baku.
        Mendukung encoding utf-8, utf-8-sig, latin-1, dan cp1252.
        Mendukung delimiter koma (,), titik-koma (;), dan tab (\t).
        Melakukan upsert (tambah baru atau perbarui jika ID Pelanggan sudah ada).
        """
        # Deteksi encoding
        decoded_text = None
        for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
            try:
                decoded_text = file_content.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if not decoded_text:
            raise ValueError("Format berkas CSV tidak dapat dibaca (encoding tidak didukung).")

        # Deteksi delimiter (koma, titik koma, atau tab)
        first_line = decoded_text.splitlines()[0] if decoded_text.splitlines() else ""
        if "\t" in first_line and first_line.count("\t") > first_line.count(",") and first_line.count("\t") > first_line.count(";"):
            delimiter = "\t"
        elif ";" in first_line and first_line.count(";") > first_line.count(","):
            delimiter = ";"
        else:
            delimiter = ","

        reader = csv.reader(io.StringIO(decoded_text), delimiter=delimiter)
        rows = list(reader)

        if not rows:
            raise ValueError("File CSV kosong.")

        # Ambil header dan normalisasi huruf kecil
        raw_headers = [h.strip().lower() for h in rows[0]]
        
        # Helper fungsi pemetaan kolom pintar (Exact & Substring matching)
        def get_col_index(candidates: List[str]) -> int:
            for idx, h in enumerate(raw_headers):
                clean_h = re.sub(r'[^a-z0-9]', '', h)
                for c in candidates:
                    clean_c = re.sub(r'[^a-z0-9]', '', c)
                    if clean_h == clean_c:
                        return idx
            for idx, h in enumerate(raw_headers):
                for c in candidates:
                    if c in h:
                        return idx
            return -1

        idx_id = get_col_index(["id pelanggan", "id_pelanggan", "id", "customer_id", "cid", "client_id", "no pelanggan"])
        idx_nama = get_col_index(["nama", "name", "nama pelanggan", "customer", "customer name", "client", "klien", "user", "member", "nama klien"])
        idx_alamat = get_col_index(["alamat", "address", "lokasi", "addr", "domisili", "tempat"])
        idx_hp = get_col_index(["no hp", "no_hp", "telepon", "whatsapp", "wa", "phone", "mobile", "kontak", "handphone"])
        idx_pop = get_col_index(["pop", "server", "olt", "wilayah", "area", "site", "lokasi server", "node"])
        idx_ip = get_col_index(["ip router", "ip_router", "ip ont", "ip modem", "ip address", "alamat ip", "ip", "host", "onu ip"])
        idx_paket = get_col_index(["paket", "bandwidth", "kecepatan", "profile", "speed", "plan", "service", "bw"])
        idx_mac = get_col_index(["mac", "mac address", "mac_address", "hardware", "mac ont"])
        idx_redaman = get_col_index(["redaman", "rx", "baseline", "rx power", "rx_power", "dbm", "sinyal", "optical power"])
        idx_wifi = get_col_index(["nama wifi", "ssid", "wifi", "wifi name"])
        idx_pass_wifi = get_col_index(["password wifi", "pass wifi", "wpa", "sandi wifi", "key wifi"])
        idx_modem = get_col_index(["jenis modem", "tipe modem", "modem", "ont", "tipe ont", "device", "model", "onu"])
        idx_user_admin = get_col_index(["user admin", "username admin", "user ont", "modem user", "username"])
        idx_pass_admin = get_col_index(["pass admin", "password admin", "pass ont", "modem pass", "password"])

        # Fallback Cerdas: Jika kolom IP belum terdeteksi dari header, scan nilai baris pertama untuk pola IPv4
        if idx_ip == -1 and len(rows) > 1:
            for r in rows[1:4]:
                for c_idx, cell in enumerate(r):
                    if re.search(r'\b(?:\d{1,3}\.){3}\d{1,3}\b', cell.strip()):
                        idx_ip = c_idx
                        break
                if idx_ip != -1:
                    break

        # Fallback Cerdas: Jika kolom Nama belum terdeteksi, ambil kolom teks pertama yang bukan IP/ID
        if idx_nama == -1:
            for c_idx, h in enumerate(raw_headers):
                if c_idx != idx_ip and c_idx != idx_id:
                    idx_nama = c_idx
                    break

        if idx_ip == -1:
            raise ValueError("Tidak dapat mendeteksi kolom 'IP Router/Modem' pada berkas CSV Anda. Pastikan terdapat kolom IP Address (contoh: 10.10.x.x).")

        imported_count = 0
        updated_count = 0
        failed_count = 0
        errors = []

        for row_idx, row in enumerate(rows[1:], start=2):
            if not row or all(not cell.strip() for cell in row):
                continue

            def get_val(idx: int) -> str:
                if 0 <= idx < len(row):
                    return row[idx].strip()
                return ""

            ip_router = get_val(idx_ip)
            nama = get_val(idx_nama) if idx_nama != -1 else ""

            if not ip_router:
                failed_count += 1
                errors.append(f"Baris {row_idx}: Alamat IP Router kosong")
                continue

            if not nama:
                nama = f"Pelanggan {ip_router}"

            # Parsing ID Pelanggan
            id_pelanggan = get_val(idx_id)
            if not id_pelanggan:
                id_pelanggan = f"CUST-{uuid.uuid4().hex[:8].upper()}"

            # Parsing Redaman Baseline
            raw_redaman = get_val(idx_redaman)
            redaman_baseline = None
            if raw_redaman:
                try:
                    cleaned_redaman = raw_redaman.replace(",", ".").replace("dBm", "").strip()
                    redaman_baseline = float(cleaned_redaman)
                except ValueError:
                    redaman_baseline = None

            pop = get_val(idx_pop) or "Server Cabang"
            alamat = get_val(idx_alamat) or None
            no_hp = get_val(idx_hp) or None
            paket = get_val(idx_paket) or None
            mac_address = get_val(idx_mac) or None
            nama_wifi = get_val(idx_wifi) or None
            password_wifi = get_val(idx_pass_wifi) or None
            jenis_modem = get_val(idx_modem) or "GM220-S"
            user_admin = get_val(idx_user_admin) or "admin"
            pass_admin = get_val(idx_pass_admin) or None

            # Cek apakah pelanggan sudah ada berdasarkan id_pelanggan atau ip_router
            existing = db.query(Pelanggan).filter(
                or_(Pelanggan.id_pelanggan == id_pelanggan, Pelanggan.ip_router == ip_router)
            ).first()

            if existing:
                existing.nama = nama
                if alamat: existing.alamat = alamat
                if no_hp: existing.no_hp = no_hp
                if pop: existing.pop = pop
                existing.ip_router = ip_router
                if paket: existing.paket = paket
                if jenis_modem: existing.jenis_modem = jenis_modem
                if mac_address: existing.mac_address = mac_address
                if redaman_baseline is not None: existing.redaman_baseline = redaman_baseline
                if nama_wifi: existing.nama_wifi = nama_wifi
                if password_wifi: existing.password_wifi = password_wifi
                if user_admin: existing.user_admin = user_admin
                if pass_admin: existing.pass_admin = pass_admin
                existing.is_active = True
                updated_count += 1
            else:
                new_cust = Pelanggan(
                    id_pelanggan=id_pelanggan,
                    nama=nama,
                    alamat=alamat,
                    no_hp=no_hp,
                    pop=pop,
                    ip_router=ip_router,
                    paket=paket,
                    jenis_modem=jenis_modem,
                    mac_address=mac_address,
                    redaman_baseline=redaman_baseline,
                    nama_wifi=nama_wifi,
                    password_wifi=password_wifi,
                    user_admin=user_admin,
                    pass_admin=pass_admin,
                    snmp_community="public",
                    is_active=True
                )
                db.add(new_cust)
                imported_count += 1

        db.commit()

        return {
            "total_processed": imported_count + updated_count + failed_count,
            "imported": imported_count,
            "updated": updated_count,
            "failed": failed_count,
            "errors": errors[:5]
        }

    @staticmethod
    def generate_excel_export(db: Session) -> bytes:
        """
        Menghasilkan file Excel (.xlsx) data seluruh pelanggan aktif
        lengkap dengan status performa dan redaman terakhir menggunakan styling profesional.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data Pelanggan ONT"
        ws.views.sheetView[0].showGridLines = True

        # Palet Warna Styling
        header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # Dark Indigo
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        # Border tipis
        thin_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)

        headers = [
            "ID Pelanggan (Wajib)", "Nama Pelanggan (Wajib)", "Alamat", "No. HP",
            "POP / OLT", "IP Router / ONT (Wajib)", "Paket Bandwidth", "Tipe Modem",
            "MAC Address", "Redaman Baseline (dBm)", "Nama WiFi", "Password WiFi",
            "User Admin", "Pass Admin",
            "Status Kredensial", "Rx Power Terakhir (dBm)", "Status Koneksi", "Pengecekan Terakhir"
        ]

        # Tulis Header
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
            # Adjust column width
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 25
        ws.row_dimensions[1].height = 28

        # Query Seluruh Data Pelanggan Aktif
        customers = db.query(Pelanggan).filter(Pelanggan.is_active == True).order_by(Pelanggan.id.asc()).all()

        for idx, c in enumerate(customers, start=1):
            latest_log = (
                db.query(LogPerformaONT)
                .filter(LogPerformaONT.id_pelanggan == c.id_pelanggan)
                .order_by(LogPerformaONT.waktu_cek.desc())
                .first()
            )
            rx_val = float(latest_log.rx_power) if (latest_log and latest_log.rx_power is not None) else (float(c.redaman_baseline) if c.redaman_baseline else "-")
            status_val = latest_log.status_koneksi if latest_log else "NORMAL"
            last_check = latest_log.waktu_cek.strftime("%d/%m/%Y %H:%M") if latest_log else "-"

            row_data = [
                c.id_pelanggan,
                c.nama,
                c.alamat or "-",
                c.no_hp or "-",
                c.pop or "-",
                c.ip_router,
                c.paket or "-",
                c.jenis_modem or "-",
                c.mac_address or "-",
                float(c.redaman_baseline) if c.redaman_baseline else "-",
                c.nama_wifi or "-",
                c.password_wifi or "-",
                c.user_admin or "-",
                c.pass_admin or "-",
                c.status_kredensial or "UNTESTED",
                rx_val,
                status_val,
                last_check
            ]

            row_idx = idx + 1
            ws.append(row_data)
            row_fill = zebra_fill if idx % 2 == 0 else white_fill

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.fill = row_fill
                cell.font = Font(name="Arial", size=9)
                cell.border = cell_border

                # Center align untuk kolom tertentu
                if col_num in [1, 2, 5, 8, 9, 10, 11, 12, 13, 14, 15, 16]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20

        # Auto adjust lebar kolom
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def generate_excel_template() -> bytes:
        """
        Menghasilkan file template Excel (.xlsx) kosong dengan format kolom standar.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template Import Pelanggan"
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        thin_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)

        headers = [
            "ID Pelanggan (Wajib)", "Nama Pelanggan (Wajib)", "Alamat", "No. HP",
            "POP / OLT", "IP Router / ONT (Wajib)", "Paket Bandwidth", "Tipe Modem",
            "MAC Address", "Redaman Baseline (dBm)", "Nama WiFi", "Password WiFi",
            "User Admin", "Pass Admin"
        ]

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 25
        ws.row_dimensions[1].height = 28

        # Tambahkan satu baris sampel sebagai contoh
        sample_row = [
            "P-001", "Budi Santoso", "Jl. Merdeka No 1", "081234567890",
            "Server Pusat", "192.168.1.100", "20Mbps", "GM220-S",
            "AA:BB:CC:DD:EE:FF", "-23.5", "Budi_WiFi", "password123",
            "admin", "admin123"
        ]
        ws.append(sample_row)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def generate_csv_template() -> str:
        """Menghasilkan isi file template CSV untuk diunduh pengguna"""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "ID Pelanggan", "Nama", "Alamat", "No HP", "POP", "IP Router",
            "Paket", "MAC Address", "Redaman", "Nama Wifi", "Password wifi",
            "Jenis Modem", "USER ADMIN", "PASS ADMIN"
        ])
        writer.writerow([
            "CUST-001", "Contoh Pelanggan 1", "Jl. Merdeka No. 10", "08123456789",
            "Server Cabang", "10.10.2.71", "10MB CAB", "44:22:95:6B:4B:C0",
            "-21.5", "WIFI-CUST01", "rahasia123", "GM220-S", "admin", "admin123"
        ])
        writer.writerow([
            "CUST-002", "Contoh Pelanggan 2", "Klero Rt 02", "08571234567",
            "Mini Klero", "10.10.2.15", "30MB Residential", "FC:8E:5B:28:4D:21",
            "-22.0", "WIFI-CUST02", "sandiwifi", "GM220-S XPON", "admin", "admin123"
        ])
        return output.getvalue()

